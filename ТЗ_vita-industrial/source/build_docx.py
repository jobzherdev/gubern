# -*- coding: utf-8 -*-
import sys, os, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MEDIA = os.path.join(ROOT, 'media')
OUT = os.path.join(ROOT, 'out', 'TZ_vita-industrial_workshops.docx')
os.makedirs(os.path.dirname(OUT), exist_ok=True)

doc = Document()

# ---------- base styles ----------
st = doc.styles['Normal']
st.font.name = 'Arial'; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12

def styleh(name, size, color, before, after):
    s = doc.styles[name]
    s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=True
    s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    s.paragraph_format.keep_with_next=True
styleh('Heading 1',16,'1F3A5F',18,8)
styleh('Heading 2',13,'2E5A88',14,6)
styleh('Heading 3',11,'333333',10,5)

sec = doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
for m in ('top_margin','bottom_margin','left_margin','right_margin'): setattr(sec,m,Cm(2))
CW = Cm(17)   # content width

# footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run('vita-industrial.ru · ТЗ на страницу «Купить производственное помещение в СПб» · ')
r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('777777')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE')
fp._p.append(fld)

# ---------- helpers ----------
def p(text='', bold=False, size=10.5, italic=False, color=None, after=6, before=0, align=None):
    par = doc.add_paragraph()
    par.paragraph_format.space_after=Pt(after); par.paragraph_format.space_before=Pt(before)
    if align: par.alignment=align
    if text:
        run=par.add_run(text); run.bold=bold; run.italic=italic; run.font.size=Pt(size)
        if color: run.font.color.rgb=RGBColor.from_string(color)
    return par

def bullets(items, size=10.5):
    for it in items:
        par=doc.add_paragraph(style='List Bullet')
        par.paragraph_format.space_after=Pt(3)
        par.add_run(it).font.size=Pt(size)

def numbered(items, size=10.5):
    for it in items:
        par=doc.add_paragraph(style='List Number')
        par.paragraph_format.space_after=Pt(3)
        par.add_run(it).font.size=Pt(size)

def shade(cell, hexc):
    tcPr=cell._tc.get_or_add_tcPr()
    sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),hexc)
    tcPr.append(sh)

def fixed_layout(t, widths):
    tblPr=t._tbl.tblPr
    for tag in ('w:tblW','w:tblLayout'):
        for e in tblPr.findall(qn(tag)): tblPr.remove(e)
    tw=OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(widths)*567))); tw.set(qn('w:type'),'dxa')
    tblPr.append(tw)
    tl=OxmlElement('w:tblLayout'); tl.set(qn('w:type'),'fixed'); tblPr.append(tl)
    grid=t._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc,w in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(int(w*567)))

def table(headers, rows, widths=None, size=9.5, head_fill='1F3A5F', zebra=True):
    ncol=len(headers) if headers else len(rows[0])
    if widths is None: widths=[17.0/ncol]*ncol
    t=doc.add_table(rows=0, cols=ncol)
    t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if headers:
        cells=t.add_row().cells
        for i,h in enumerate(headers):
            cells[i].text=''
            par=cells[i].paragraphs[0]; par.paragraph_format.space_after=Pt(2)
            run=par.add_run(h); run.bold=True; run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string('FFFFFF')
            shade(cells[i], head_fill)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=''
            par=cells[i].paragraphs[0]; par.paragraph_format.space_after=Pt(2)
            texts = val if isinstance(val,list) else [val]
            for j,txt in enumerate(texts):
                if j: par=cells[i].add_paragraph(); par.paragraph_format.space_after=Pt(2)
                bold = txt.startswith('**')
                run=par.add_run(txt.replace('**',''))
                run.font.size=Pt(size); run.bold=bold
            if zebra and ri%2==1: shade(cells[i],'F4F6F9')
    for r_ in t.rows:
        for i,w in enumerate(widths):
            r_.cells[i].width=Cm(w)
    fixed_layout(t, widths)
    return t

def kv_table(rows, w1=4.6, size=9.5):
    return table(None, rows, widths=[w1, 17-w1], size=size, zebra=False)

def pagetext(lines):
    """Блок «ТЕКСТ ДЛЯ СТРАНИЦЫ»"""
    par=doc.add_paragraph(); par.paragraph_format.space_before=Pt(6); par.paragraph_format.space_after=Pt(2)
    run=par.add_run('ТЕКСТ ДЛЯ СТРАНИЦЫ'); run.bold=True; run.font.size=Pt(9.5)
    run.font.color.rgb=RGBColor.from_string('2E5A88')
    for ln in lines:
        if isinstance(ln,tuple):
            kind,val=ln
            if kind=='h': p(val, bold=True, size=10.5, after=3, before=4)
            elif kind=='b': bullets(val)
            elif kind=='n': numbered(val)
        else:
            p(ln, after=5)

def pic(fname, caption, width=17.0):
    path=os.path.join(MEDIA,fname)
    par=doc.add_paragraph(); par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before=Pt(8); par.paragraph_format.space_after=Pt(2)
    par.add_run().add_picture(path, width=Cm(width))
    cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after=Pt(12)
    r=cp.add_run(caption); r.font.size=Pt(8.5); r.italic=True
    r.font.color.rgb=RGBColor.from_string('555555')

FIG=[0]
def fig(fname, text, width=17.0):
    FIG[0]+=1
    pic(fname, f'Рис. {FIG[0]}. {text}', width)

# ================= TITLE =================
tp=doc.add_paragraph(); tp.paragraph_format.space_after=Pt(2)
r=tp.add_run('Техническое задание на страницу «Купить производственное помещение в Санкт-Петербурге»')
r.bold=True; r.font.size=Pt(17); r.font.color.rgb=RGBColor.from_string('1F3A5F')
sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(14)
r=sp.add_run('vita-industrial.ru/workshops · индустриальный парк «Вита», Санкт-Петербург, Шушары · SEO-продвижение')
r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string('666666')

p('Документ описывает, что должно быть на странице https://vita-industrial.ru/workshops, чтобы она собирала трафик по коммерческим запросам о покупке производственных помещений в Санкт-Петербурге. '
  'Для каждого блока указаны заголовок, тег, формат вёрстки и готовый текст для страницы. Шапка и футер сайта используются текущие, без изменений. '
  'В конце документа — макет наполнения страницы: общий вид и все блоки по порядку.', after=10)

# ================= 1. META =================
doc.add_heading('1. Метатеги и служебные элементы', level=1)
p('Заполняются в настройках страницы Tilda: «Настройки страницы» → «Главное» (адрес страницы) и «SEO» (Title, Description, Open Graph).', size=9.5, color='555555')
table(['Элемент','Значение','Длина'],
      [[a,b,c] for a,b,c in C.META],
      widths=[4.2,11.2,1.6])

# ================= 2. AUDIT =================
doc.add_heading('2. Что мешает текущей странице ранжироваться', level=1)
p('Страница уже существует и проиндексирована, поэтому её не переносят на новый адрес, а перестраивают. Ниже — восемь проблем, которые нужно закрыть; они же объясняют, зачем нужны новые блоки из раздела 5.', after=8)
table(['Проблема','Как сейчас','Что делаем'],
      [[a,b,c] for a,b,c in C.AUDIT], widths=[4.4,6.6,6.0])
fig('now_top.jpg','Текущая страница vita-industrial.ru/workshops: H1 «Производственные цеха для развития вашего бизнеса» без продвигаемых слов, на странице нет ни одного тега H2')

# ================= 3. QUERIES =================
doc.add_heading('3. Продвигаемые запросы и их распределение по блокам', level=1)
p('WS — базовая частотность, !WS — точная. Запросы распределены так, чтобы каждый блок отвечал на свою группу и текст оставался читаемым. Точные формы («производственные помещения купить спб») используются один раз, в тексте вводного или итогового блока.', after=8)
table(['Фраза','WS','!WS','Где используется'],
      [[a,b,c,d] for a,b,c,d in C.QUERIES], widths=[6.2,1.5,1.5,7.8])
p('Правила употребления: «купить» и «продажа» обязательны в H1, в блоках 5, 13, 16 и 20. Слово «склад» на этой странице используется только в связке «производственно-складское помещение» и в ссылке на раздел «Склады»: складская лексика уводит страницу от продвигаемых запросов.', size=9.5, after=10)

# ================= 4. COMPETITORS =================
doc.add_heading('4. Конкуренты из топа: что берём в структуру', level=1)
p('Разобраны пять страниц, которые находятся в топе по продвигаемым запросам. Четыре из них — агрегаторы и агентства, у которых нет собственных объектов; пятая — индустриальный парк с похожей бизнес-моделью. Отсюда вывод: наша страница выигрывает конкретикой (реальные площади, цены, сети, сроки сделки), которой у агрегаторов нет.', after=8)
for i,(url,h1,what,take) in enumerate(C.COMPETITORS,1):
    doc.add_heading(f'4.{i}. {url}', level=3)
    kv_table([['**H1 конкурента**',h1],['**Что на странице**',what],['**Что берём**',take]], w1=4.0, size=9.5)
    if i<=4:
        fig(f'comp_k{i}.jpg', f'{url} — первый экран страницы из топа')
    else:
        p('Скриншот не приводится: сайт недоступен из среды подготовки документа, анализ сделан по выдаче и кэшу.', size=9, italic=True, color='777777')
doc.add_heading('Вывод для структуры', level=3)
bullets([
 'Ни у одного конкурента в топе нет цен на конкретные объекты с площадями — наш блок 13 закрывает главный вопрос пользователя прямо на странице.',
 'У всех сильных страниц есть большой связный текст с H2/H3 про выбор, районы, сделку и сравнение покупки с арендой — это блоки 5, 10, 14, 16.',
 'FAQ с разметкой есть только у части конкурентов и в усечённом виде — блок 20 из 14 вопросов даёт преимущество в расширенном сниппете.',
 'Агрегаторы выигрывают количеством объявлений, поэтому нам нужен каталог с фильтром по площади и живым статусом свободных блоков (блоки 6 и 8).'])

# ================= 5. STRUCTURE =================
doc.add_heading('5. Структура страницы', level=1)
p('Порядок блоков построен по логике покупателя: что за объект → какие помещения есть → что внутри → где находится → кому подходит → сколько стоит → почему купить выгоднее → как проходит сделка → доверие → вопросы → заявка. Существующие блоки сохраняются на своих местах, новые встроены там, где их ждёт посетитель.', after=8)
rows=[]
for b in C.BLOCKS:
    rows.append([str(b['n']), b['name'], b['head'], b['tag'], b['status']])
table(['№','Блок','Заголовок','Тег','В шаблоне'], rows, widths=[1.0,3.6,7.2,1.4,3.8], size=9)
p('CTA-строки. После блоков 6, 13 и 17 размещается тёмная полоса с текстом и кнопкой «Оставить заявку». Посетитель не должен пролистать больше двух экранов без кнопки обращения.', size=9.5, after=4)
p('Скрытые блоки. Блоки 18 и 19 верстаются сразу, но показываются после согласования кейсов с резидентами и появления отзывов.', size=9.5, after=10)

# ================= 6. BLOCKS =================
doc.add_heading('6. Блоки страницы: содержание и тексты', level=1)

def block_head(b):
    doc.add_heading(f'Блок {b["n"]}. {b["name"]}', level=2)
    kv_table([['**Заголовок и тег**', f'{b["tag"]}: {b["head"]}' if b['tag'] not in ('—',) else 'Без заголовка'],
              ['**Статус**', b['status']],
              ['**Формат**', b['note']]], w1=3.6)

for b in C.BLOCKS:
    n=b['n']; L=b['layout']; d=b.get('data')
    block_head(b)
    if L=='header':
        p('Используем текущую шапку: логотип, меню «О компании / Объекты / Ход строительства / Документация / Цены / Контакты», телефон +7 (812) 574-47-47, мессенджеры, кнопка «Заказать звонок». В меню «Объекты» проверить, что ссылка на «Цеха» ведёт на /workshops.', after=6)
    elif L=='crumbs':
        pagetext(['Главная / Объекты / Производственные помещения'])
        p('Последний элемент — текст без ссылки. Разметка BreadcrumbList (раздел 7).', size=9.5, after=6)
    elif L=='hero':
        pagetext([('h','H1: '+d['h1']), d['sub'], ('h','Строка с ценой'), d['price'],
                  ('h','Кнопки'), 'Кнопка 1: «Подобрать помещение» — открывает форму заявки.',
                  'Кнопка 2: «Смотреть цены и планировки» — якорь на блок 13.',
                  ('h','Четыре плашки'), ('b', d['chips'])])
    elif L=='numbers':
        pagetext([('h','H2: '+b['head']), ('b',[f'{a} — {x}' for a,x in d])])
    elif L=='text2col':
        pagetext([('h','H2: '+b['head'])]+list(d['paras'])+[('h','Изображение справа'), d['img']])
        p('Объём текста — около 1 900 знаков. Это главный текстовый блок страницы, правки согласовывать с SEO-специалистом.', size=9.5, italic=True, after=6)
    elif L=='catalog':
        pagetext([('h','H2: '+b['head']), d['lead']])
        table(['Карточка (H3)','Текст карточки','Покупка','Аренда'],
              [[t, txt, p1, p2] for t,_,txt,p1,p2 in d['items']], widths=[3.6,8.4,2.5,2.5], size=9)
        p(d['note'], size=9.5, after=4)
        bullets(['Заголовки карточек меняем со «Складское помещение N м²» на «Производственный цех N м²».',
                 'В каждой карточке — кнопка «Скачать планировку PDF» (файл с именем вида plan-tseh-750.pdf).',
                 'Alt у планировок: «План производственного помещения 750 м², индустриальный парк «Вита», Шушары».'], size=9.5)
    elif L=='spectable':
        pagetext([('h','H2: '+b['head'])])
        table(['Параметр','Значение'], [[a,x] for a,x in d], widths=[5.5,11.5], size=9.5)
    elif L=='genplan':
        pagetext([('h','H2: '+b['head']), d,
                  'Подпись под схемой: «На схеме — 36 зданий парка. Зелёным отмечены свободные производственные помещения, серым — забронированные и проданные. Статус обновляется еженедельно, последнее обновление: [дата]».'])
    elif L=='accordion':
        pagetext([('h','H2: '+b['head'])])
        table(['Пункт аккордеона','Текст'], [[a,x] for a,x in d], widths=[4.6,12.4], size=9.5)
    elif L=='location':
        pagetext([('h','H2: '+b['head']), ('h','Плашки у карты'), ('b', d['points']), ('h','Текст под картой')]+list(d['paras']))
    elif L=='gallery':
        pagetext([('h','H2: '+b['head']), ('h','Состав галереи'), ('b', d),
                  'Подпись под галереей: «Строительство идёт очередями, актуальные фотографии — в разделе «Ход строительства»».'])
    elif L=='cards6':
        pagetext([('h','H2: '+b['head'])])
        table(['Карточка','Текст'], [[a,x] for a,x in d], widths=[5.0,12.0], size=9.5)
        p('Под карточками сохраняем текущий блок «Не нашли свой бизнес?» с формой, заменив слово «склады» на «помещения».', size=9.5)
    elif L=='pricetable':
        pagetext([('h','H2: '+b['head'])])
        table(['Помещение','Общая площадь','Цена за м²','Стоимость покупки','Аренда'], [list(r) for r in d['rows']],
              widths=[4.6,2.8,2.4,3.6,3.6], size=9)
        p('Под таблицей: «Прайс актуален на [дата]. Цены указаны без НДС / с НДС — уточнить у бухгалтерии». Кнопка «Скачать прайс PDF».', size=9.5, after=6)
        pagetext(list(d['paras'])+[('h','H3: Условия оплаты')])
        table(['Вариант','Условия'], [[a,x] for a,x in d['terms']], widths=[4.4,12.6], size=9.5)
    elif L=='compare':
        pagetext([('h','H2: '+b['head'])])
        table(d['header'], [list(r) for r in d['rows']], widths=[3.4,7.0,6.6], size=9.5)
        p(d['after'], after=6)
        p('Расчёт окупаемости: 105 000 ₽/м² ÷ (800 ₽/м² × 12 месяцев) ≈ 10,9 года. Перед публикацией перепроверить с коммерческим отделом и при изменении прайса пересчитать.', size=9, italic=True, color='777777')
    elif L=='cards3':
        pagetext([('h','H2: '+b['head'])])
        table(['Карточка','Текст'], [[a,x] for a,x in d], widths=[4.4,12.6], size=9.5)
    elif L=='steps':
        pagetext([('h','H2: '+b['head'])])
        table(['Шаг','Название','Текст'], [[str(i+1),t,x] for i,(t,x) in enumerate(d)], widths=[1.2,4.0,11.8], size=9.5)
        p('Под лентой: «Средний срок от первого звонка до получения ключей — около трёх недель. Менеджер: +7 (812) 574-47-47».', size=9.5)
    elif L=='docs':
        pagetext([('h','H2: '+b['head']), ('b', d)])
        p('Справа — плитки со ссылками на PDF из раздела «Документация»: разрешение на ввод объекта, образец выписки ЕГРН, технические условия на сети, типовой договор купли-продажи, планировки блоков. Файлы отдавать с говорящими именами, не «doc1.pdf».', size=9.5)
    elif L=='cases':
        pagetext([('h','H2: '+b['head'])])
        table(['Отрасль','Площадь','Что разместили'], [[a,s,x] for a,s,x in d], widths=[4.4,2.4,10.2], size=9.5)
        p('Перед публикацией согласовать с резидентами упоминание отрасли и фото. Если согласования нет — блок остаётся скрытым, на ранжирование это не влияет.', size=9.5, italic=True)
    elif L=='reviews':
        pagetext([('h','H2: '+b['head']),
                  'До появления собственных отзывов выводим виджет Яндекс.Карт с карточкой организации.',
                  'По мере появления добавляем 5–10 текстовых отзывов: имя, компания, отрасль, площадь купленного блока, дата.',
                  'Отзывы обязательно выводить текстом в HTML, а не только скриптом виджета: иначе поисковые системы их не увидят.'])
    elif L=='faq':
        pagetext([('h','H2: '+b['head'])])
        table(['Вопрос','Ответ'], [[q,a] for q,a in d], widths=[5.0,12.0], size=9.5)
        p('Первые три ответа открыты по умолчанию, остальные свёрнуты. Под списком — кнопка «Задать вопрос менеджеру». Разметка FAQPage ставится на тот же текст, который видит пользователь.', size=9.5)
    elif L=='form':
        pagetext([('h','H2: '+b['head']), d['text'],
                  ('h','Поля формы'), ('b', d['fields']+['Чекбокс «Нужна рассрочка»','Чекбокс согласия с политикой конфиденциальности']),
                  ('h','Кнопка'), d['btn']])
        p('Исправить опечатки в текущих текстах сайта: «Оставьте заявкук» → «Оставьте заявку», «по всем вопрсам» → «по всем вопросам», «под ваш запроос» → «под ваш запрос».', size=9.5, italic=True)
    elif L=='seotail':
        pagetext([('h','H2: '+b['head'])]+list(d['paras']))
        table(['Анкор ссылки','Адрес'], [[t,u] for t,u in d['links']], widths=[8.5,8.5], size=9.5)
    elif L=='footer':
        p('Используем текущий футер: логотип, телефон, почта, адрес офиса, навигация, объекты, политика конфиденциальности. Это общий блок сайта, правка на этой странице изменит его везде.', after=6)

# ================= 7. MARKUP =================
doc.add_heading('7. Микроразметка', level=1)
table(['Тип разметки','Где и что размечаем'], [[a,x] for a,x in C.MARKUP], widths=[4.6,12.4], size=9.5)
p('Разметку ставить в формате JSON-LD в теле страницы. После публикации проверить валидатором Яндекса и Google Rich Results Test.', size=9.5)

# ================= 8. TECH =================
doc.add_heading('8. Технические требования и проверка после публикации', level=1)
table(['Требование','Что сделать'], [[a,x] for a,x in C.TECH], widths=[4.2,12.8], size=9.5)

# ================= APPENDIX =================
doc.add_page_break()
doc.add_heading('Приложение. Макет наполнения страницы', level=1)
p('Макет показывает страницу /workshops сверху вниз: все 23 блока, CTA-строки, шапку и футер сайта — с заголовками и текстами из раздела 6, в цветах и шрифтах vita-industrial.ru. '
  'Фотографии, планировки, карта и схема парка показаны заглушками: их берут из текущих материалов сайта.', after=8)
table(['Метка на макете','Значение'],
      [['Оранжевая','Блок есть в текущем шаблоне страницы, наполняем и дорабатываем по ТЗ'],
       ['Фиолетовая','Новый блок, в шаблоне его нет'],
       ['Серая','Блок свёрстан, но скрыт до появления материалов (кейсы, отзывы)'],
       ['Синяя','Общий блок сайта: шапка и футер, не меняем'],
       ['Тёмно-зелёная полоса','CTA-строка с кнопкой заявки между блоками']], widths=[4.6,12.4], size=9.5)

fig('fig_overview.jpg','Общий вид страницы: читать сверху вниз, колонки 1, 2, 3')
fig('fig_mobile.jpg','Первый экран на мобильном (390 px): H1, подзаголовок, цена и кнопка заявки видны без прокрутки', width=7.2)

parts=json.load(open(os.path.join(ROOT,'parts.json')))
total=len(parts)
def short(lbls):
    out=[]
    for l in lbls:
        l=l.strip()
        if l.startswith('CTA'): out.append('CTA-строка')
        else:
            t=l.replace('Блок ','').split(' · ')
            out.append(f'{t[0]} — {t[1]}' if len(t)>1 else t[0])
    return '; '.join(out)
for i,pt in enumerate(parts,1):
    fig(f'fig_part{i:02d}.jpg', f'Макет страницы, часть {i} из {total}: {short(pt["labels"])}')

doc.save(OUT)
print('saved', OUT, os.path.getsize(OUT)//1024, 'KB')
